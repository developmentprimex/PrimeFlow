import asyncio
from datetime import date
from io import BytesIO

from docx import Document
from openpyxl import load_workbook

from app.services.one_h_ga_summary import build_ga_summary, summary_html
from app.services.primeflow_report import build_report_document, render_html, render_plain_text, render_docx, render_png
from app.services.tomorrow_print_report import build_today_print_report, _build_print_report


DAY = date(2026, 10, 15)


def payload():
    return {"departments": [{"id": "dev", "code": "DEV"}, {"id": "gd", "code": "GD"}, {"id": "pcm", "code": "PCM"}],
            "users": [
                {"id": "fg", "full_name": "Fiona Gashi", "department_id": "gd", "weekly_planner_sort_order": 0},
                {"id": "lh", "full_name": "Lina Hoxha", "department_id": "dev", "weekly_planner_sort_order": 3},
                {"id": "ef", "full_name": "Era Fana", "department_id": "dev", "weekly_planner_sort_order": 0},
                {"id": "ra", "full_name": "Rina Aliu", "department_id": "dev", "weekly_planner_sort_order": 1},
                {"id": "ep", "full_name": "Ema Pasha", "department_id": "pcm", "weekly_planner_sort_order": 0},
            ]}


def task(key, marker=None, **changes):
    return {"id": key, "task_id": key, "title": "EF/GA: Work", "person": "Era Fana",
            "assignees": ["Era Fana"], "date": DAY.isoformat(), "status": "TODO",
            "one_h_marker": marker, "one_h_report_slot": "10:00", "finish_period": "AM", **changes}


def test_single_row_follows_planner_order_and_deduplicates_overlapping_categories():
    personal = task("personal", "FLAG")
    m2 = task("m2", "M2_M3")
    shared = task("shared", "F", assignees=["Era Fana", "Fiona Gashi"])
    entries = [(personal, True), (personal, True), (personal, False), (m2, False), (m2, False), (shared, False),
               (task("not-ga", title="EF/KA: Other personal"), True),
               (task("unmarked", title="EF: Other work"), False)]
    rows = build_ga_summary(payload(), entries)
    assert [row.person for row in rows] == ["EF", "FG"]
    assert rows[0].model_dump() == {"person": "EF", "department": "ZHV", "personal": 1,
                                    "ga": 1, "m2": 1, "physical": 1, "total": 3}
    assert rows[1].physical == rows[1].total == 1
    assert all(row.total > 0 for row in rows)
    markup = summary_html(rows)
    assert markup.count("<tr>") == 1
    assert markup.count("<td ") == 2
    assert markup.index("<strong>EF</strong>") < markup.index("<strong>FG</strong>")
    assert "(ZHV)" not in markup and "(DSG)" not in markup and "NR DET ME GA" not in markup
    assert "<strong>FG</strong> - F:1" in markup
    assert "<strong>EF</strong> - F:1 GA:1 P:1 M2:1" in markup
    assert "P:0" not in markup and "<br>" not in markup


def test_inactive_and_admin_owners_are_excluded():
    data = payload()
    data["users"].extend([
        {"id": "ga", "full_name": "Gane Arifaj", "role": "ADMIN", "department_id": "dev"},
        {"id": "old", "full_name": "Old Person", "is_active": False, "department_id": "dev"},
    ])
    rows = build_ga_summary(data, [
        (task("admin", "F", person="Gane Arifaj", assignees=["Gane Arifaj"]), False),
        (task("inactive", "F", person="Old Person", assignees=["Old Person"]), False)])
    assert rows == []
    assert summary_html(rows) == ""


def test_rap_1h_summary_is_present_in_html_text_word_and_png():
    data = payload()
    data["items"] = {"oneH": [task("flag", "FLAG"), task("m2", "M2")], "personal": [task("p")]}
    document = build_report_document(data, DAY, "10:00")
    assert document.ga_task_summary[0].total == 3
    markup = render_html(document)
    assert 'data-ga-task-summary="true"' in markup
    assert "EF - GA:1 P:1 M2:1" in render_plain_text(document)
    word = Document(BytesIO(render_docx(document)))
    assert any("EF - GA:1 P:1 M2:1" in cell.text for table in word.tables for row in table.rows for cell in row.cells)
    assert render_png(document).startswith(b"\x89PNG")


def test_shtypi_summary_has_one_row_and_matches_excel_and_word():
    data = payload()
    data["items"] = {"oneH": [task("flag", "FLAG"), task("m2", "M2")], "personal": [task("p")]}
    report = asyncio.run(_build_print_report(DAY, payload=data, include_attachment=True, include_docx=True,
                                            include_png=True, report_day_label="SOT"))
    assert report["ga_task_summary"][0]["total"] == 3
    summary = report["html"].split('data-ga-task-summary="true"', 1)[1].split("</table>", 1)[0]
    assert summary.count("<tr>") == 1
    assert "EF - GA:1 P:1 M2:1" in report["plain_text"]
    excel = load_workbook(BytesIO(report["attachments"][0][1]))
    assert any("EF - GA:1 P:1 M2:1" in str(cell.value) for row in excel.active for cell in row)
    attachment = next(item for item in report["attachments"] if item[0].endswith(".docx"))
    word = Document(BytesIO(attachment[1]))
    assert any("EF - GA:1 P:1 M2:1" in cell.text for table in word.tables for row in table.rows for cell in row.cells)
