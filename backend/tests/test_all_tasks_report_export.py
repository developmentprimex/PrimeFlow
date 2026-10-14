import io
import unittest
from types import SimpleNamespace

from openpyxl import load_workbook

from app.api.routers.exports import AllTasksReportExportIn, export_all_tasks_report_xlsx
from app.models.enums import UserRole


class TestAllTasksReportExport(unittest.IsolatedAsyncioTestCase):
    async def test_finance_export_removes_markup_and_preserves_text_and_breaks(self):
        for frequency in (None, "3M"):
            with self.subTest(frequency=frequency):
                payload = AllTasksReportExportIn(
                    title="ALL FINANCE TASK REPORT",
                    rows=[{
                        "title": "[[added]]HS:[[/added]] [[ADDED]]BORXH HC[[/ADDED]]",
                        "description": "<p>Muaji &amp; raporti</p>[[done]]Drafti[[/done]]<br>GA<br/>MAX<BR />25",
                        "details": "&lt;br&gt;[[added]]REGJ: 0[[/added]]\r\n<div>PATH: 0</div>",
                        "frequency": frequency,
                        "department": "FIN",
                        "status": "TODO",
                        "userInitials": "HS",
                    }, {}],
                )
                response = await export_all_tasks_report_xlsx(
                    payload,
                    user=SimpleNamespace(role=UserRole.ADMIN, full_name="HS", username="hs"),
                )
                content = b"".join([chunk async for chunk in response.body_iterator])
                worksheet = load_workbook(io.BytesIO(content)).active
                columns = {cell.value: cell.column for cell in worksheet[4]}
                expected = {
                    "TITULLI": "HS: BORXH HC",
                    "PERSHKRIMI": "Muaji & raporti\nDrafti\nGA\nMAX\n25",
                    "REGJ/PATH/CHECKLISTA/TRAINING/GROUP": "REGJ: 0\nPATH: 0",
                    "DEP": "FIN",
                    "STATUS": "TODO",
                    "USER": "HS",
                }
                for header, value in expected.items():
                    cell = worksheet.cell(5, columns[header])
                    self.assertEqual(cell.value, value)
                    self.assertTrue(cell.alignment.wrap_text)
                self.assertEqual(worksheet.cell(6, columns["TITULLI"]).value, "-")
                self.assertIsNone(worksheet.cell(6, columns["PERSHKRIMI"]).value)
                self.assertEqual("F" in columns, frequency is not None)
