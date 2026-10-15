import asyncio
import io
import uuid
from datetime import date, datetime
from types import SimpleNamespace
from unittest.mock import AsyncMock, MagicMock, patch

import pytest
from docx import Document
from openpyxl import load_workbook
from pydantic import ValidationError

from app.models.enums import UserRole
from app.api.routers import m3_reporting_points as api
from app.services import m3_reporting_points as report_service
from app.services import m3_reporting_points_attachments as exports
from app.services import m3_unheld_meetings as meetings
from app.services.primeflow_report import report_timezone

DAY = date(2026, 10, 14)


def meeting(kind="internal", hour=10, **values):
    data = dict(id=uuid.uuid4(), title="Team sync", meeting_type=kind,
                starts_at=datetime(2026, 10, 14, hour, tzinfo=report_timezone()), ends_at=None,
                recurrence_type=None, recurrence_days_of_week=None, recurrence_days_of_month=None,
                department_id=uuid.uuid4(), calendar_sync_status=None, calendar_imported=False)
    data.update(values)
    return SimpleNamespace(**data)


def test_only_today_canceled_int_ext_meetings_are_selected():
    internal, external, held, canceled = meeting(), meeting("external", 15), meeting(), meeting(hour=9)
    hidden = meeting(calendar_sync_status="cancelled")
    leave = meeting("external", title="PV annual leave", calendar_imported=True, calendar_categories=["PV"])
    yesterday = meeting(starts_at=datetime(2026, 10, 13, 10, tzinfo=report_timezone()))
    rows = meetings.select_unheld_meetings([external, held, internal, canceled, hidden, leave, yesterday],
                                         {internal.id: "canceled", external.id: "cancelled", held.id: "held", canceled.id: "canceled",
                                          hidden.id: "canceled", leave.id: "canceled", yesterday.id: "canceled"}, DAY)
    assert [row["meeting_id"] for row in rows] == [str(canceled.id), str(internal.id), str(external.id)]
    assert [row["meeting_type"] for row in rows] == ["INT", "INT", "EXT"]
    assert rows[0]["status_label"] == "Anuluar"
    assert rows[2]["time"] == "15:00"


def test_recurring_and_multiday_meetings_use_the_report_date_occurrence():
    recurring = meeting(recurrence_type="weekly", recurrence_days_of_week=[DAY.weekday()],
                        starts_at=datetime(2026, 9, 30, 10, tzinfo=report_timezone()))
    other_day = meeting(recurrence_type="weekly", recurrence_days_of_week=[(DAY.weekday() + 1) % 7])
    multiday = meeting("external", starts_at=datetime(2026, 10, 13, 9, tzinfo=report_timezone()),
                       ends_at=datetime(2026, 10, 15, 12, tzinfo=report_timezone()))
    rows = meetings.select_unheld_meetings([recurring, other_day, multiday],
                                         {recurring.id: "canceled", other_day.id: "canceled", multiday.id: "canceled"}, DAY)
    assert {row["meeting_id"] for row in rows} == {str(recurring.id), str(multiday.id)}
    assert all(row["occurrence_date"] == DAY.isoformat() for row in rows)
    assert meetings.select_unheld_meetings([recurring], {recurring.id: "held"}, DAY) == []


def test_builder_checks_daily_status_and_reuses_manual_meeting_palette():
    current = meeting()
    def result(items):
        value = MagicMock()
        value.scalars.return_value.all.return_value = items
        return value
    db = AsyncMock()
    db.execute.side_effect = [result([current]), result([SimpleNamespace(meeting_id=current.id, status="canceled")])]
    rows = asyncio.run(meetings.build_unheld_meetings(db, DAY))
    assert rows[0]["origin"] == "MANUAL"
    assert rows[0]["fill"] == "#FFF4CC"
    assert rows[0]["highlighted"]
    status_query = db.execute.call_args_list[1].args[0]
    assert DAY in status_query.compile().params.values()
    assert db.execute.await_count == 2


def test_planning_comment_keys_are_validated_and_other_keys_are_rejected():
    key = f"meeting_plan:{uuid.uuid4()}"
    assert api.ManualAnswersPayload(manual_answers={key: "Nesër 10:00"}).manual_answers[key] == "Nesër 10:00"
    for invalid in ("meeting_plan:", "meeting_plan:not-a-uuid", "unknown"):
        with pytest.raises(ValidationError):
            api.ManualAnswersPayload(manual_answers={invalid: "text"})
    with pytest.raises(ValidationError):
        api.ManualAnswersPayload(manual_answers={key: "x" * 10001})


@pytest.mark.parametrize("status", ["DRAFT", "SENT"])
def test_staff_can_save_meeting_comment_before_capture_or_after_send(status):
    key = f"meeting_plan:{uuid.uuid4()}"
    row = SimpleNamespace(id=uuid.uuid4(), report_date=DAY, manual_answers={"underload": "Ruaj"},
                          data={}, realization={"percent": 62}, realization_captured_at=None,
                          generated_at=None, status=status, sent_at=None, last_error=None)
    db = SimpleNamespace(get=AsyncMock(return_value=row), refresh=AsyncMock(), commit=AsyncMock(), add=MagicMock())
    with patch.object(api, "locked_report", new=AsyncMock(return_value=row)):
        payload = asyncio.run(api.save_answers(row.id, api.ManualAnswersPayload(manual_answers={key: "Nesër 10:00"}),
                                               db, SimpleNamespace(id=uuid.uuid4(), role=UserRole.STAFF, full_name="Staff")))
    assert payload["manual_answers"] == {"underload": "Ruaj", key: "Nesër 10:00"}
    assert payload["realization"] == {"percent": 62}
    db.commit.assert_awaited_once()


def test_meeting_title_and_planning_comment_reach_email_word_excel_and_png_blocks():
    source = meeting(title="INT <script> test")
    row = meetings.select_unheld_meetings([source], {source.id: "canceled"}, DAY)[0]
    key = f"meeting_plan:{row['meeting_id']}"
    report = {"report_date": DAY.isoformat(), "manual_answers": {key: "Nesër 10:00 <ora>"},
              "data": {"unheld_meetings": [row]}, "realization": None}
    html = report_service.render_html(report)
    assert "KUR PLANIFIKOHEN?" in html
    assert "INT &lt;script&gt; test" in html
    assert "Nesër 10:00 &lt;ora&gt;" in html
    assert "TAK EXTERNE" in html and "TAK INTERNE" in html
    assert "#FFF4CC" in html and "#2563eb" in html
    assert "MBAJTUR?" in html and "KOHA" in html
    assert "Nesër 10:00 <ora>" in report_service.render_plain_text(report)
    blocks = exports.export_blocks(report)
    table = next(block for block in blocks if block.get("columns") == meetings.MEETING_COLUMNS)
    assert table["rows"][0][-1]["text"] == "Nesër 10:00 <ora>"
    word = Document(io.BytesIO(exports.render_docx(report)))
    text = "\n".join(cell.text for table in word.tables for line in table.rows for cell in line.cells)
    assert "KUR PLANIFIKOHEN?" in text and "Nesër 10:00 <ora>" in text
    from app.services.reporting_points_excel import render_xlsx
    sheet = load_workbook(io.BytesIO(render_xlsx(blocks, sheet_name="M3"))).active
    values = [cell.value for line in sheet for cell in line]
    assert "KUR PLANIFIKOHEN?" in values and "Nesër 10:00 <ora>" in values
    assert exports.render_png(report).startswith(b"\x89PNG")


def test_linked_int_uses_common_view_para_pas_links_and_external_palette():
    external = meeting("external", title="Weekly meeting", calendar_imported=True, calendar_categories=["Weekly"])
    internal = meeting(title="Internal after", paired_external_meeting_id=external.id)
    before = meeting(title="Internal before", pre_external_meeting_id=external.id)
    rows = meetings.select_unheld_meetings([internal, before, external],
                                         {internal.id: "canceled", before.id: "canceled", external.id: "canceled"}, DAY)
    by_id = {row["meeting_id"]: row for row in rows}
    assert by_id[str(internal.id)]["meeting_relation"] == "PAS"
    assert by_id[str(before.id)]["meeting_relation"] == "PARA"
    assert by_id[str(external.id)]["meeting_relation"] is None
    assert all(row["origin"] == "CAL" and row["fill"] == "#C9A98A" for row in rows)
    groups = meetings.meeting_groups({"report_date": DAY.isoformat(), "data": {"unheld_meetings": rows}})
    assert [(label, len(items)) for label, items in groups] == [("TAK EXTERNE", 1), ("TAK INTERNE", 2)]
    for label, items in groups:
        values = meetings.meeting_values(items[0], 1)
        assert values[:3] == ["1", "10:00", "✕"]
        assert "[CAL]" in values[3]
        if label == "TAK INTERNE":
            assert values[3].startswith("[PAS]") or values[3].startswith("[PARA]")
        assert len(values) == 5


def test_saved_data_also_filters_unmarked_held_and_other_day_meetings():
    rows = [{"meeting_id": str(uuid.uuid4()), "meeting_type": "INT", "occurrence_date": day, "status": status}
            for day, status in [(DAY.isoformat(), "canceled"), (DAY.isoformat(), "held"),
                                (DAY.isoformat(), ""), ("2026-10-13", "canceled")]]
    report = {"report_date": DAY.isoformat(), "data": {"unheld_meetings": rows}}
    assert len(meetings.meeting_rows(report)) == 1


def test_staff_cannot_change_manager_manual_questions_through_comment_save():
    db = SimpleNamespace(get=AsyncMock())
    user = SimpleNamespace(id=uuid.uuid4(), role=UserRole.STAFF, full_name="Staff")
    payload = api.ManualAnswersPayload(manual_answers={"underload": "Unauthorized"})
    with pytest.raises(api.HTTPException) as error:
        asyncio.run(api.save_answers(uuid.uuid4(), payload, db, user))
    assert error.value.status_code == 403
    db.get.assert_not_awaited()


def test_saved_meeting_appearance_restores_linked_colour_without_changing_snapshot():
    external = meeting("external", calendar_imported=True, calendar_categories=["Weekly"])
    internal = meeting(title="Current title", paired_external_meeting_id=external.id)
    old_row = {"meeting_id": str(internal.id), "meeting_type": "INT", "occurrence_date": DAY.isoformat(),
               "title": "Saved title", "time": "14:20", "status": "canceled"}
    data = {"unheld_meetings": [old_row], "untouched": [{"task_id": "saved"}]}
    def result(items):
        value = MagicMock()
        value.scalars.return_value.all.return_value = items
        return value
    db = AsyncMock()
    db.execute.side_effect = [result([internal]), result([external])]
    updated = asyncio.run(meetings.add_saved_meeting_appearance(db, data))
    restored = updated["unheld_meetings"][0]
    assert restored["fill"] == "#C9A98A" and restored["origin"] == "CAL"
    assert restored["highlighted"]
    assert restored["meeting_relation"] == "PAS"
    assert restored["title"] == "Saved title" and restored["status"] == "canceled"
    assert restored["time"] == "14:20"
    assert updated["untouched"] == data["untouched"]
    assert "fill" not in old_row
    assert db.execute.await_count == 2
    db.commit.assert_not_awaited()


def test_reports_with_saved_appearance_need_no_extra_queries():
    data = {"unheld_meetings": [{"meeting_id": str(uuid.uuid4()), "fill": "#C9A98A", "origin": "", "highlighted": False, "meeting_relation": None}]}
    db = AsyncMock()
    assert asyncio.run(meetings.add_saved_meeting_appearance(db, data)) is data
    db.execute.assert_not_awaited()


def test_current_report_payload_restores_colour_and_keeps_saved_comment_and_generation_time():
    source = meeting("external", calendar_imported=True, calendar_categories=["Weekly"])
    key = f"meeting_plan:{source.id}"
    row = SimpleNamespace(id=uuid.uuid4(), report_date=DAY, manual_answers={key: "Nesër 10:00"},
                          data={"unheld_meetings": [{"meeting_id": str(source.id), "title": source.title,
                                "meeting_type": "EXT", "occurrence_date": DAY.isoformat(), "status": "canceled"}]},
                          realization=None, realization_captured_at=None, generated_at=source.starts_at,
                          status="DRAFT", sent_at=None, last_error=None)
    result = MagicMock()
    result.scalars.return_value.all.return_value = [source]
    db = AsyncMock()
    db.execute.return_value = result
    payload = asyncio.run(report_service.current_report_payload(db, row))
    assert payload["data"]["unheld_meetings"][0]["fill"] == "#C9A98A"
    assert payload["manual_answers"][key] == "Nesër 10:00"
    assert payload["generated_at"] == source.starts_at.isoformat()
    assert "fill" not in row.data["unheld_meetings"][0]
