import io
import unittest
import uuid
from types import SimpleNamespace
from unittest.mock import AsyncMock, patch

from fastapi import FastAPI
from fastapi.testclient import TestClient
from openpyxl import load_workbook

from app.api.routers.exports import export_repetitive_tasks_xlsx, router
from app.models.department import Department
from app.models.enums import UserRole
from app.models.system_task_template import SystemTaskTemplate
from app.models.system_task_template_assignee_slot import SystemTaskTemplateAssigneeSlot
from app.models.task import Task
from app.models.user import User
from app.services.repetitive_task_report import repetitive_task_report_rows


class Result:
    def __init__(self, rows):
        self.rows = rows

    def scalars(self):
        return self

    def all(self):
        return self.rows


def database(*results):
    return SimpleNamespace(execute=AsyncMock(side_effect=[Result(rows) for rows in results]))


class TestRepetitiveTaskReport(unittest.IsolatedAsyncioTestCase):
    async def test_sources_merge_people_and_departments_once_per_logical_task(self):
        dev, pcm = uuid.uuid4(), uuid.uuid4()
        alice, bob = uuid.uuid4(), uuid.uuid4()
        template_id = uuid.uuid4()
        note_id, plan_id, group_id = uuid.uuid4(), uuid.uuid4(), uuid.uuid4()
        template = SystemTaskTemplate(
            id=template_id, title="System work", department_id=dev,
            assignee_ids=[uuid.uuid4()], scope="ALL",
        )
        slots = [SystemTaskTemplateAssigneeSlot(template_id=template_id, primary_user_id=person)
                 for person in (alice, bob)]
        tasks = []
        for origin, value in (("ga_note_origin_id", note_id), ("plan_note_origin_id", plan_id),
                              ("fast_task_group_id", group_id)):
            for person, dep in ((alice, dev), (bob, pcm)):
                tasks.append(Task(id=uuid.uuid4(), title="Shared work", assigned_to=person,
                                  department_id=dep, **{origin: value}))
        # Same title without shared provenance is a distinct task.
        separate = Task(id=uuid.uuid4(), title="Shared work", assigned_to=alice, department_id=dev)
        tasks.append(separate)
        users = [User(id=alice, full_name="Alice Example", department_id=dev),
                 User(id=bob, full_name="Bob Example", department_id=pcm)]
        departments = [Department(id=dev, code="DEV"), Department(id=pcm, code="PCM")]
        db = database([template], tasks, slots, [(separate.id, alice)], users, departments)
        rows = await repetitive_task_report_rows(db)
        self.assertEqual(len(rows), 5)
        self.assertEqual(rows[0], ["System work", "Alice Example, Bob Example", "DEV / PCM", "System"])
        self.assertEqual(rows.count(["Shared work", "Alice Example, Bob Example", "DEV / PCM", "Jo system"]), 3)
        self.assertIn(["Shared work", "Alice Example", "DEV", "Jo system"], rows)
        queries = [str(call.args[0]) for call in db.execute.call_args_list]
        self.assertIn("system_task_templates.is_repetitive IS true", queries[0])
        self.assertIn("system_task_templates.approval_status =", queries[0])
        self.assertIn("tasks.is_repetitive IS true", queries[1])
        self.assertIn("tasks.is_active IS true", queries[1])
        self.assertIn("tasks.system_template_origin_id IS NULL", queries[1])
        self.assertNotIn("tasks.status =", queries[1])

    async def test_empty_report_and_unassigned_template(self):
        self.assertEqual(await repetitive_task_report_rows(database([], [])), [])
        template = SystemTaskTemplate(id=uuid.uuid4(), title="All staff", scope="ALL")
        self.assertEqual(await repetitive_task_report_rows(database([template], [], [])),
                         [["All staff", "All", "ALL", "System"]])

    async def test_legacy_assignees_and_multiple_task_assignees(self):
        alice, bob = uuid.uuid4(), uuid.uuid4()
        template = SystemTaskTemplate(id=uuid.uuid4(), title="Legacy", assignee_ids=[alice], default_assignee_id=bob)
        task = Task(id=uuid.uuid4(), title="Project work")
        users = [User(id=alice, full_name="Alice"), User(id=bob, full_name="Bob")]
        db = database([template], [task], [], [(task.id, alice), (task.id, bob)], users)
        rows = await repetitive_task_report_rows(db)
        self.assertEqual(rows, [["Legacy", "Alice, Bob", "-", "System"],
                                ["Project work", "Alice, Bob", "-", "Jo system"]])

    async def test_excel_design_numbering_text_and_empty_export(self):
        for rows in ([], [["=A1", "Alice, Bob", "DEV / PCM", "System"],
                          ["Long title " * 20, "Alice", "DEV", "Jo system"]]):
            for role in (UserRole.ADMIN, UserRole.MANAGER, UserRole.STAFF):
                with self.subTest(rows=len(rows), role=role):
                    user = SimpleNamespace(role=role, full_name="Test User", username="test")
                    with patch("app.services.repetitive_task_report.repetitive_task_report_rows",
                               new=AsyncMock(return_value=rows)):
                        response = await export_repetitive_tasks_xlsx(db=database(), user=user)
                    content = b"".join([chunk async for chunk in response.body_iterator])
                    ws = load_workbook(io.BytesIO(content)).active
                    self.assertEqual([cell.value for cell in ws[5]], ["NR", "TITULLI", "PERSONAT", "DEPARTAMENTI", "TYPE"])
                    self.assertEqual(ws["A3"].value, "DETYRA TË PËRSËRITURA")
                    self.assertEqual(ws["B5"].fill.fgColor.rgb, "00D9D9D9")
                    self.assertEqual(ws["B5"].border.bottom.style, "thin")
                    self.assertEqual(ws.auto_filter.ref, f"A5:E{5 + len(rows)}")
                    self.assertEqual(ws.freeze_panes, "C6")
                    self.assertEqual(ws.page_setup.fitToHeight, 0)
                    if rows:
                        self.assertEqual(list(ws.values)[5:], [(index, *row) for index, row in enumerate(rows, 1)])
                        self.assertEqual(ws["B6"].data_type, "s")
                        self.assertTrue(ws["B7"].alignment.wrap_text)
                        self.assertGreater(ws.row_dimensions[7].height, 30)

    def test_requires_authentication(self):
        app = FastAPI()
        app.include_router(router, prefix="/exports")
        with TestClient(app) as client:
            self.assertEqual(client.get("/exports/repetitive-tasks.xlsx").status_code, 401)
