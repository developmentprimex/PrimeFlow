import asyncio
import re
import uuid
from datetime import date, datetime, timezone
from io import BytesIO
from types import SimpleNamespace
from unittest.mock import AsyncMock, MagicMock

import pytest
from fastapi import Request, Response
from openpyxl import load_workbook
from sqlalchemy import create_engine, text

from app.api.routers import common_view
from app.models.enums import UserRole
from app.models.meeting import Meeting
from app.services.tomorrow_print_report import (
    _docx_table_attachment,
    _excel_table_attachment,
    _meeting_rows,
    build_today_print_report,
    build_tomorrow_print_report,
)


def meeting_items():
    return {
        bucket: [
            {"title": f"{bucket} {day} {name}", "date": day, "time": time,
             "occurrenceStatus": status, "recurrence_type": "weekly"}
            for day in ("2026-10-28", "2026-10-29")
            for name, time, status in (
                ("canceled", "08:00", "canceled"),
                ("held", "10:00", "held"),
                ("pending", "09:00", ""),
            )
        ] for bucket in ("external", "internal")
    }


@pytest.mark.parametrize("tomorrow", [False, True])
def test_both_reports_show_status_and_only_canceled_meetings_move_last(tomorrow):
    builder = build_tomorrow_print_report if tomorrow else build_today_print_report
    report = asyncio.run(builder(
        date(2026, 10, 27) if tomorrow else date(2026, 10, 28),
        payload={"items": meeting_items()},
    ))
    for output in (report["html"], report["content_html"]):
        assert output.count('data-meeting-status="held"') == 4
        assert output.count('data-meeting-status="canceled"') == 4
        assert ">Mbajtur</span>" in output
        assert ">Anuluar X</span>" in output
        assert output.count("background-image:linear-gradient") == 4
        pending_cells = [body for body in re.findall(
            r'<td data-meeting-time="true"[^>]*>(.*?)</td>', output,
        ) if body.startswith("09:00")]
        assert len(pending_cells) == 4
        assert all("data-meeting-status" not in body for body in pending_cells)
        for day in ("2026-10-28", "2026-10-29"):
            for bucket in ("external", "internal"):
                titles = [f"{bucket} {day} {name}" for name in ("pending", "held", "canceled")]
                assert output.index(titles[0]) < output.index(titles[1]) < output.index(titles[2])
    assert "[Mbajtur]" in report["plain_text"]
    assert "[Anuluar X]" in report["plain_text"]
    assert "[]" not in report["plain_text"]


def test_exports_keep_status_cancellation_mark_and_order():
    from docx import Document

    rows = [(label, values, False) for label, values in _meeting_rows(meeting_items(), date(2026, 10, 28))]
    sections = [(date(2026, 10, 28), "SOT", rows)]
    _, content, _ = _excel_table_attachment([], [], date(2026, 10, 28), meeting_sections=sections)
    sheet = load_workbook(BytesIO(content)).active
    title_cells = [row[5] for row in sheet if str(row[5].value or "").startswith(tuple(f"{i}. " for i in range(1, 7)))]
    assert [str(cell.value).split(". ", 1)[1] for cell in title_cells] == [
        f"{prefix}{bucket} 2026-10-28 {name}"
        for bucket in ("external", "internal")
        for prefix, name in (("", "pending"), ("[Mbajtur] ", "held"), ("[Anuluar X] ", "canceled"))
    ]
    assert [cell.font.strike is True for cell in title_cells] == [False, False, True] * 2
    _, content, _ = _docx_table_attachment([], date(2026, 10, 28), meeting_sections=sections)
    document = Document(BytesIO(content))
    meeting_rows = [row for table in document.tables for row in table.rows if "2026-10-28" in row.cells[-1].text]
    assert len(meeting_rows) == 6
    assert "[Mbajtur]" in meeting_rows[1].cells[-1].text
    assert "[Anuluar X]" in meeting_rows[2].cells[-1].text
    assert all(run.font.strike for run in meeting_rows[2].cells[-1].paragraphs[0].runs)


def test_common_view_captures_status_per_occurrence_for_recurring_and_one_off_meetings(monkeypatch):
    recurring = Meeting(
        id=uuid.uuid4(), title="Recurring", meeting_type="internal",
        starts_at=datetime(2026, 10, 28, 8, tzinfo=timezone.utc),
        recurrence_type="weekly", recurrence_days_of_week=[2, 3, 4],
    )
    single = Meeting(
        id=uuid.uuid4(), title="One off", meeting_type="external", recurrence_type="none",
        starts_at=datetime(2026, 10, 28, 8, tzinfo=timezone.utc),
    )
    statuses = [
        SimpleNamespace(meeting_id=recurring.id, occurrence_date=date(2026, 10, 28), status="held"),
        SimpleNamespace(meeting_id=recurring.id, occurrence_date=date(2026, 10, 29), status="canceled"),
        SimpleNamespace(meeting_id=single.id, occurrence_date=date(2026, 10, 28), status="canceled"),
    ]

    def result(values):
        value = MagicMock()
        value.scalars.return_value.all.return_value = values
        value.all.return_value = values
        return value

    db = AsyncMock()
    db.execute.side_effect = [result([]), result([]), result([recurring, single]), result(statuses), result([])]
    monkeypatch.setattr(common_view, "_compute_etag", AsyncMock(return_value="test-meeting-status"))
    payload = asyncio.run(common_view.get_common_view(
        request=Request({"type": "http", "headers": [(b"cache-control", b"no-store")]}),
        response=Response(), week_start=date(2026, 10, 26), include="meetings",
        department_id=None, include_all_departments=True, db=db,
        user=SimpleNamespace(role=UserRole.ADMIN),
    ))
    assert [(item["date"], item["occurrenceStatus"]) for item in payload.items.internal] == [
        ("2026-10-28", "held"), ("2026-10-29", "canceled"), ("2026-10-30", ""),
    ]
    assert payload.items.external[0]["occurrenceStatus"] == "canceled"


def test_status_changes_reset_and_deletion_invalidate_common_view_cache():
    engine = create_engine("sqlite://")
    with engine.begin() as connection:
        connection.execute(text("CREATE TABLE departments (created_at DATETIME)"))
        connection.execute(text("CREATE TABLE meetings (updated_at DATETIME, department_id TEXT)"))
        connection.execute(text("CREATE TABLE meeting_occurrence_statuses (occurrence_date DATE, updated_at DATETIME)"))
        db = SimpleNamespace(execute=AsyncMock(side_effect=lambda query: connection.execute(query)))

        def etag():
            return asyncio.run(common_view._compute_etag(
                db, date(2026, 10, 26), date(2026, 11, 1), ["meetings"],
                SimpleNamespace(role=UserRole.ADMIN), None, True, False,
                [date(2026, 10, 28)], date(2026, 10, 28), date(2026, 10, 28),
            ))

        before = etag()
        connection.execute(text("INSERT INTO meeting_occurrence_statuses VALUES ('2026-10-28', '2026-10-28 09:00:00')"))
        added = etag()
        connection.execute(text("UPDATE meeting_occurrence_statuses SET updated_at='2026-10-28 10:00:00'"))
        changed = etag()
        connection.execute(text("DELETE FROM meeting_occurrence_statuses"))
        assert len({before, added, changed}) == 3
        assert etag() == before
    engine.dispose()
