import asyncio
import uuid
from datetime import date, datetime, timezone
from types import SimpleNamespace

import pytest
from pydantic import ValidationError

from app.services import m1_reporting_points as service
from app.models.enums import CommonCategory, GaNoteStatus
from app.api.routers.m1_reporting_points import ManualAnswersPayload


DAY = date(2026, 10, 13)


def ga_task(**changes):
    return SimpleNamespace(**{
        "id": uuid.uuid4(), "title": "EF: Detyre", "status": "TODO", "progress_percentage": 0,
        "created_at": DAY, "start_date": DAY, "due_date": DAY, "finish_period": "AM",
        "completed_at": None, "is_deadline_important": False, "one_h_marker": None, **changes,
    })


def test_ga_tasks_keep_personal_symbols_and_m2_delivery_in_separate_tables():
    origin = uuid.uuid4()
    tasks = [ga_task(one_h_marker="F"), ga_task(one_h_marker="FLAG"),
             ga_task(one_h_marker="M2", finish_period="PM"),
             ga_task(one_h_marker="M2_M3", finish_period="PM"),
             ga_task(one_h_marker="M2", is_personal=True, ga_note_origin_id=origin)]
    excluded = [ga_task(one_h_marker="M3"), ga_task(one_h_marker="M2", due_date=date(2026, 10, 12)),
                ga_task(one_h_marker="F", finish_period="PM")]
    sections = service.select_ga_tasks(tasks + excluded, [], DAY,
                                      personal_titles={origin: "EF/GA: Personal"})
    rows = sections["ga_progress"]
    assert {row["task_id"] for row in rows} == {str(tasks[index].id) for index in (0, 1, 4)}
    assert {row["task_id"] for row in sections["delivery"]} == {str(tasks[index].id) for index in (2, 3, 4)}
    assert len(rows) == 3
    assert next(row for row in rows if row["task_id"] == str(tasks[-1].id))["title"] == "EF/GA: Personal"
    report = {"report_date": DAY.isoformat(), "data": sections, "manual_answers": {}}
    tables = [block for block in service.export_blocks(report) if block["kind"] == "table"]
    assert len(tables) == 2
    table = tables[0]
    assert [column[0] for column in table["columns"]] == [
        "nr", "assignees", "department", "project", "am_pm", "task_type", "marker", "title"]
    assert len(table["rows"]) == 3
    assert len(tables[1]["rows"]) == 3
    assert "M2/3" in service.render_html(report)
    assert "EF/GA: Personal" in service.render_plain_text(report)


def at(day, hour, minute=0):
    return datetime(day.year, day.month, day.day, hour, minute, tzinfo=service.report_timezone())


@pytest.mark.parametrize("hour,minute", [(8, 15), (10, 0), (16, 45)])
def test_generation_windows_include_boundaries_and_exclude_future_creations(hour, minute):
    yesterday = date(2026, 10, 12)
    now = at(DAY, hour, minute)
    after = now.replace(microsecond=1)
    notes = [
        SimpleNamespace(id="before", content="Too early", created_at=at(yesterday, 16, 29)),
        SimpleNamespace(id="start", content="Start", created_at=at(yesterday, 16, 30)),
        SimpleNamespace(id="end", content="End", created_at=now),
        SimpleNamespace(id="after", content="Too late", created_at=after),
    ]
    tasks = [
        SimpleNamespace(id="old", title="Yesterday", created_at=at(yesterday, 23), status="TODO"),
        SimpleNamespace(id="today", title="Today", created_at=at(DAY, 0), status="TODO"),
        SimpleNamespace(id="cutoff", title="Cutoff", created_at=now, status="TODO"),
        SimpleNamespace(id="later", title="Later", created_at=after, status="TODO"),
    ]
    data = service.select_sections([], notes, [], tasks, DAY, now=now)
    assert [row["content"] for row in data["notes"]] == ["End", "Start"]
    assert [row["title"] for row in data["tasks"]] == ["Today", "Cutoff"]


def test_new_tasks_include_meetings_owner_department_period_and_type():
    task = SimpleNamespace(id="new", title="Detyra\nPërshkrimi", created_at=at(DAY, 8), status="TODO",
                           assigned_to="ada", department_id="other", finish_period="PM", project_id=None,
                           system_template_origin_id=None, system_task_slot_id=None,
                           is_bllok=False, is_r1=False, is_1h_report=False, is_personal=False)
    users = {"ada": SimpleNamespace(full_name="Ada Test", username="ada")}
    data = service.select_sections([], [], [], [task], DAY, now=at(DAY, 10), users=users,
                                           departments={"other": "DEV"}, assignees={"new": {"ada"}})
    row = data["tasks"][0]
    assert (row["assignees"], row["department"], row["am_pm"], row["task_type"]) == ("AT", "DEV", "PM", "FT")


def test_created_tasks_split_by_local_start_day_independently_of_due_date():
    tasks = [
        ga_task(title="Start today", created_at=at(DAY, 8),
                start_date=datetime(2026, 10, 12, 22, 30, tzinfo=timezone.utc), due_date=date(2026, 10, 20)),
        ga_task(title="Start tomorrow", created_at=at(DAY, 8), start_date=date(2026, 10, 14), due_date=DAY),
        ga_task(title="Start yesterday", created_at=at(DAY, 8), start_date=date(2026, 10, 12)),
        ga_task(title="No start", created_at=at(DAY, 8), start_date=None),
    ]
    data = service.select_sections([], [], [], tasks, DAY, now=at(DAY, 9))
    assert [row["start_date"] for row in data["tasks"]] == [DAY.isoformat(), "2026-10-14", "2026-10-12", None]
    report = {"report_date": DAY.isoformat(), "data": data, "manual_answers": {}}
    blocks = service.export_blocks(report)
    headings = [block["text"] for block in blocks if block["kind"] == "text" and block.get("level") == 3]
    assert headings == ["1. TË KRIJUARA SOT ME START DATE SOT", "2. TË KRIJUARA SOT ME START DATE NË DITË TJETËR"]
    tables = [block for block in blocks if block["kind"] == "table"]
    assert [[row[5]["text"] for row in table["rows"]] for table in tables] == [
        ["Start today"], ["Start tomorrow", "Start yesterday", "No start"]]
    assert [row[0]["text"] for row in tables[1]["rows"]] == ["1", "2", "3"]
    assert "Përfshihen edhe detyrat pa Start Date." in service.render_plain_text(report)


@pytest.mark.parametrize("due,label", [
    (DAY, "Sot"), (date(2026, 10, 14), "Nesër"),
    (date(2026, 10, 19), "19.10.2026"), (None, "Pa afat"),
])
def test_new_tasks_split_due_dates_and_marker_into_export_columns(due, label):
    task = ga_task(created_at=at(DAY, 8), due_date=due, one_h_marker="FLAG")
    data = service.select_sections([], [], [], [task], DAY, now=at(DAY, 10))
    row = data["tasks"][0]
    assert row["marker"] == "GA"
    assert row["planned_for"] == label
    report = {"report_date": DAY.isoformat(), "data": data}
    cells = service._table_rows(report, "tasks")[0]
    assert len(cells) == len(service.TABLE_COLUMNS["tasks"])
    exported = dict(zip([column[0] for column in service.TABLE_COLUMNS["tasks"]], cells))
    assert exported["marker"]["text"] == "GA"
    assert exported["planned_for"]["text"] == label


def test_new_notes_include_only_px_notes():
    yesterday = date(2026, 10, 12)
    ga_note = SimpleNamespace(id="ga", content="PX note", created_at=at(yesterday, 17), created_by="ada")
    plan_note = SimpleNamespace(id="plan", content="JAV note", created_at=at(DAY, 8), created_by="beni")
    users = {"ada": SimpleNamespace(full_name="Ada Test"), "beni": SimpleNamespace(full_name="Beni Test")}

    data = service.select_sections([], [ga_note], [plan_note], [], DAY, now=at(DAY, 10), users=users)

    assert [(row["source"], row["author"]) for row in data["notes"]] == [("PX Notes", "Ada Test")]


def test_sent_report_backfills_only_missing_planning_and_preserves_saved_values():
    from unittest.mock import AsyncMock, MagicMock
    today = datetime.now(service.report_timezone()).date()
    item = ga_task(due_date=today)
    saved = {"id": str(item.id), "title": "Saved title", "marker": "F"}
    existing = {"id": str(uuid.uuid4()), "planned_for": "Nesër", "marker": "M2"}
    row = SimpleNamespace(report_date=today, data={"tasks": [saved, existing]},
                          sent_at=at(today, 8), manual_answers={"reorganization": "Jo"})
    result = MagicMock()
    result.scalars.return_value.all.return_value = [item]
    db = AsyncMock()
    db.execute.return_value = result
    assert asyncio.run(service.backfill_saved_task_planning(db, row))
    assert row.data["tasks"] == [{**saved, "planned_for": "Sot", "start_date": DAY.isoformat()},
                                 {**existing, "start_date": None}]
    assert row.sent_at == at(today, 8)
    assert row.manual_answers == {"reorganization": "Jo"}
    assert not asyncio.run(service.backfill_saved_task_planning(db, row))
    db.execute.assert_awaited_once()
    row.report_date = date(2020, 1, 1)
    row.data = {"tasks": [saved]}
    assert not asyncio.run(service.backfill_saved_task_planning(db, row))
    db.execute.assert_awaited_once()


def test_notes_converted_to_tasks_are_excluded_from_new_reports():
    notes = [
        SimpleNamespace(id="open", content="Keep", created_at=at(DAY, 8), is_converted_to_task=False),
        SimpleNamespace(id="converted", content="Hide", created_at=at(DAY, 8), is_converted_to_task=True),
    ]

    data = service.select_sections([], notes, [], [], DAY, now=at(DAY, 10))

    assert [note["id"] for note in data["notes"]] == ["open"]


def test_closed_notes_are_excluded_from_new_reports():
    notes = [
        SimpleNamespace(id="open", content="Keep", created_at=at(DAY, 8), status=GaNoteStatus.OPEN),
        SimpleNamespace(id="closed-enum", content="Hide enum", created_at=at(DAY, 8), status=GaNoteStatus.CLOSED),
        SimpleNamespace(id="closed-string", content="Hide string", created_at=at(DAY, 8), status="CLOSED"),
    ]

    data = service.select_sections([], notes, [], [], DAY, now=at(DAY, 10))

    assert [note["id"] for note in data["notes"]] == ["open"]


@pytest.mark.parametrize("origin", ["from_mira", "from_mira_desktop"])
@pytest.mark.parametrize("shared", [False, True])
def test_mira_notes_are_excluded_from_new_reports(origin, shared):
    notes = [
        SimpleNamespace(id="open", content="Keep", created_at=at(DAY, 8)),
        SimpleNamespace(id="mira", content="Meeting summary", created_at=at(DAY, 8),
                        mira_shared_to_px=shared, **{origin: True}),
    ]
    data = service.select_sections([], notes, [], [], DAY, now=at(DAY, 10))
    assert [note["id"] for note in data["notes"]] == ["open"]


def test_build_data_queries_only_open_unconverted_notes(monkeypatch):
    statements = []

    class Result:
        def scalars(self):
            return self

        def all(self):
            return []

    class Db:
        async def execute(self, statement):
            statements.append(statement)
            return Result()

    async def context(db, tasks):
        return tasks, [], {}, {}

    async def ga_context(db, day, cutoff):
        return {"ga_progress": []}

    monkeypatch.setattr(service, "prepare_task_context", context)
    monkeypatch.setattr(service, "build_ga_task_data", ga_context)
    asyncio.run(service.build_data(Db(), DAY, now=at(DAY, 10)))

    notes_query = statements[1]
    assert "ga_notes.status =" in str(notes_query)
    assert GaNoteStatus.OPEN in notes_query.compile().params.values()
    assert "ga_notes.is_converted_to_task IS false" in str(notes_query)
    for statement in statements:
        assert at(DAY, 10) in statement.compile().params.values()


def test_saved_report_hides_notes_closed_or_converted_after_generation_and_puts_done_tasks_last():
    stored = {"notes": [
        {"id": "converted", "source": "PX Notes", "content": "Hide"},
        {"id": "closed", "source": "PX Notes", "content": "Closed"},
        {"id": "open", "source": "PX Notes", "content": "Keep"},
    ], "tasks": [
        {"id": "done", "status": "DONE"},
        {"id": "todo", "status": "TODO"},
        {"id": "waiting", "status": "WAITING_CLIENT"},
    ]}
    row = SimpleNamespace(id="report", report_date=DAY, manual_answers={}, data=stored, status="DRAFT",
                          generated_at=None, sent_at=None, last_error=None)

    payload = service.report_payload(row, excluded_note_ids={"converted", "closed"})

    assert [note["id"] for note in payload["data"]["notes"]] == ["open"]
    assert [task["id"] for task in payload["data"]["tasks"]] == ["todo", "waiting", "done"]
    assert [note["id"] for note in stored["notes"]] == ["converted", "closed", "open"]
    assert [task["id"] for task in stored["tasks"]] == ["done", "todo", "waiting"]


def test_excluded_note_lookup_reads_only_px_notes_and_checks_closed_or_converted():
    ga_id, plan_id = uuid.uuid4(), uuid.uuid4()
    row = SimpleNamespace(data={"notes": [
        {"id": str(ga_id), "source": "PX Notes"},
        {"id": str(plan_id), "source": "PX JAV"},
    ]})

    class Result:
        def scalars(self):
            return self

        def all(self):
            return [ga_id]

    class Db:
        async def execute(self, statement):
            params = list(statement.compile().params.values())
            assert any(isinstance(value, (set, list, tuple)) and ga_id in value and plan_id not in value
                       for value in params)
            assert "is_converted_to_task" in str(statement)
            assert " OR " in str(statement)
            assert GaNoteStatus.CLOSED in params
            assert "discord_voice_meetings.saved_note_id = ga_notes.id" in str(statement)
            assert "mira_1h_events.note_id = ga_notes.id" in str(statement)
            assert "mira_desktop_sessions.note_id = ga_notes.id" in str(statement)
            return Result()

    assert asyncio.run(service.excluded_note_ids(Db(), row)) == {str(ga_id)}


def test_older_saved_reports_hide_px_jav_notes_without_changing_saved_data():
    stored = {"notes": [
        {"id": "ga", "source": "PX Notes", "content": "Keep"},
        {"id": "plan", "source": "PX JAV", "content": "Hide"},
    ], "tasks": []}
    row = SimpleNamespace(id="report", report_date=DAY, manual_answers={}, data=stored, status="DRAFT",
                          generated_at=None, sent_at=None, last_error=None)

    payload = service.report_payload(row)

    assert [note["id"] for note in payload["data"]["notes"]] == ["ga"]
    assert [note["id"] for note in stored["notes"]] == ["ga", "plan"]


def test_saved_px_notes_gain_authors_without_changing_their_content_or_order():
    data = {"notes": [
        {"id": "ga", "source": "PX Notes", "content": "First"},
        {"id": "old", "source": "PX Notes", "content": "Third", "author": "Gane Arifaj"},
    ], "tasks": [{"id": "task"}]}

    updated = service.add_saved_note_authors(data, {("PX Notes", "ga"): "Ada Test"})

    assert [row["author"] for row in updated["notes"]] == ["Ada Test", "Gane Arifaj"]
    assert [row["content"] for row in updated["notes"]] == ["First", "Third"]
    assert updated["tasks"] == data["tasks"]
    assert "author" not in data["notes"][0]


def test_new_tasks_follow_weekly_planner_order():
    def task(name, order):
        return SimpleNamespace(id=name, title=name, created_at=at(DAY, 8), assigned_to=name,
                               is_deadline_important=False, fast_task_order=None,
                               system_template_origin_id=None, system_task_slot_id=None,
                               _weekly_planner_report_sort=order)

    tasks = [task("pcm", (2, "pcm", 0, 1, "pcm")), task("dev-last", (0, "dev", 0, 2, "dev-last")),
             task("dev-first", (0, "dev", 0, 1, "dev-first"))]
    assert [item.id for item in service.sort_tasks_for_report(tasks, {}, {})] == ["dev-first", "dev-last", "pcm"]

    tasks[2].status = "DONE"
    assert [item.id for item in service.sort_tasks_for_report(tasks, {}, {})] == ["dev-last", "pcm", "dev-first"]


def test_saved_unsent_report_tasks_can_be_reordered_without_replacing_answers_or_rows():
    data = {"notes": [{"id": "note"}], "tasks": [
        {"id": "pcm", "title": "First", "status": "TODO"},
        {"id": "dev-last", "title": "Second", "status": "DONE"},
        {"id": "dev-first", "title": "Third", "status": "TODO"},
    ]}
    updated = service.reorder_saved_task_rows(data, ["dev-first", "dev-last", "pcm"])
    assert [row["id"] for row in updated["tasks"]] == ["dev-first", "dev-last", "pcm"]
    assert updated["notes"] == data["notes"]
    assert updated["tasks"][1]["status"] == "DONE"
    assert updated["task_order_version"] == 1
    assert [row["id"] for row in data["tasks"]] == ["pcm", "dev-last", "dev-first"]


def test_common_view_absences_and_delays_use_report_day_and_person():
    entries = [
        SimpleNamespace(id="a", category=CommonCategory.absences, entry_date=DAY, created_at=at(DAY, 7),
                        assigned_to_user_id="person", created_by_user_id="creator", title="Other", description="From: 08:00 - To: 12:00 Sick"),
        SimpleNamespace(id="b", category=CommonCategory.delays, entry_date=DAY, created_at=at(DAY, 7),
                        assigned_to_user_id="person", created_by_user_id="creator", title="Other", description="Start: 08:00 Until: 09:00 Traffic"),
        SimpleNamespace(id="c", category="absences", entry_date=date(2026, 10, 12), created_at=at(DAY, 7),
                        assigned_to_user_id="person", created_by_user_id="creator", title="Other", description=None),
    ]
    users = {"person": SimpleNamespace(full_name="Ada Test", username="ada")}
    data = service.select_sections(entries, [], [], [], DAY, now=at(DAY, 10), users=users)
    assert [(row["person"], row["from"], row["to"]) for row in data["absences"]] == [("Ada Test", "08:00", "12:00")]
    assert [(row["person"], row["start"], row["until"]) for row in data["delays"]] == [("Ada Test", "08:00", "09:00")]


def test_render_includes_all_six_points_and_manual_comments_escaped():
    report = {"report_date": DAY.isoformat(), "manual_answers": {"plan_changes": "Po <sot>", "reorganization": "Jo"},
              "data": {"absences": [{"id": "absence", "person": "Ada", "from": "08:00", "to": "12:00", "note": "Sëmurë"}],
                       "delays": [], "notes": [{"id": "note", "source": "PX Notes", "content": "Shënim i ri", "created_at": at(DAY, 8).isoformat()}],
                       "tasks": [{"id": "task", "title": "Detyrë e re", "status": "TODO", "created_at": at(DAY, 8).isoformat()}]}}
    html = service.render_html(report)
    assert all(f"{number}." in html for number in range(1, 7))
    assert "Po &lt;sot&gt;" in html and "Po <sot>" not in html
    assert "Jo" in service.render_plain_text(report)


def test_only_the_two_manual_points_can_be_saved():
    payload = ManualAnswersPayload(manual_answers={"plan_changes": "Ndryshim", "reorganization": "Po"})
    assert len(payload.manual_answers) == 2
    with pytest.raises(ValidationError):
        ManualAnswersPayload(manual_answers={"notes": "Nuk lejohet"})


def test_report_exports_include_the_six_points():
    report = {"report_date": DAY.isoformat(), "subject": service.subject_for(DAY),
              "manual_answers": {"plan_changes": "Po", "reorganization": "Jo"},
              "data": {"absences": [], "delays": [],
                       "notes": [{"id": "note", "source": "PX Notes", "content": "Shënim i ri", "created_at": at(DAY, 8).isoformat()}],
                       "tasks": [{"id": "task", "title": "Detyrë e re", "status": "TODO",
                                  "created_at": at(DAY, 8).isoformat(), "assignees": "AT",
                                  "department": "DEV", "am_pm": "PM", "task_type": "FT"}]}}
    attachments = service.report_attachments(report)
    assert [name.rsplit(".", 1)[-1] for name, _, _ in attachments] == ["docx", "png", "xlsx"]
    assert all(content for _, content, _ in attachments)


def test_report_tables_keep_meetings_columns_and_category_colors():
    report = {"report_date": DAY.isoformat(), "manual_answers": {"plan_changes": "Po", "reorganization": "Jo"},
              "data": {
                  "absences": [{"id": "a", "person": "Ada", "from": "08:00", "to": "12:00", "note": "Sëmurë"}],
                  "delays": [{"id": "b", "person": "Beni", "start": "08:00", "until": "09:00", "note": "Trafik"}],
                  "notes": [{"id": "n", "source": "PX Notes", "author": "Ada Test", "content": "[[added]]Shënim i ri[[/added]]", "created_at": at(DAY, 8).isoformat()}],
                  "tasks": [{"id": "t", "title": "Detyrë e re", "status": "IN_PROGRESS", "created_at": at(DAY, 8).isoformat(),
                             "assignees": "AT", "department": "DEV", "am_pm": "PM", "task_type": "FT"}],
              }}
    tables = [block for block in service.export_blocks(report) if block["kind"] == "table"]
    assert [[column[1] for column in table["columns"]] for table in tables] == [
        ["NR", "PERSONI", "ORA"],
        ["NR", "PERSONI", "ORA"],
        ["NR", "KUSH", "NOTE", "KRIJUAR"],
        ["NR", "KUSH", "DEP", "AM/PM", "LLOJI", "DETYRA", "SIMBOLI", "PËR KUR", "KRIJUAR"],
    ]
    assert [cell["text"] for cell in tables[3]["rows"][0][1:5]] == ["AT", "DEV", "PM", "FT"]
    assert [cell["text"] for cell in tables[2]["rows"][0][1:]] == ["AT", "Shënim i ri", "13.10.2026 08:00"]
    assert tables[3]["rows"][0][5]["text"] == "Detyrë e re"
    assert tables[3]["rows"][0][8]["text"] == "13.10.2026 08:00"
    assert [len(table["rows"][0]) for table in tables[:2]] == [3, 3]
    assert [table["rows"][0][0]["fill"] for table in tables] == ["#fee2e2", "#fef3c7", "#dbeafe", "#ffff00"]
    html = service.render_html(report)
    assert "background:#e2e8f0" in html
    assert "background:#fee2e2" in html
    assert "PERSONI" in html and "KUSH" in html and "BURIMI" not in html and "STATUSI" not in html
    assert "SHËNIMI" not in html and "Trafik" not in html and "Sëmurë" not in html
    assert "[[added]]" not in html and "Shënim i ri" in html


def test_task_exports_show_only_the_first_title_line():
    report = {"report_date": DAY.isoformat(), "manual_answers": {},
              "data": {"absences": [], "delays": [], "notes": [], "tasks": [
                  {"id": "t", "title": "\nOH: [[added]]08:00[[/added]] EM: CC EMAIL CLAIM\nPërshkrimi i gjatë i detyrës",
                   "status": "TODO", "created_at": at(DAY, 8).isoformat()},
              ]}}
    task_table = next(block for block in service.export_blocks(report) if block["kind"] == "table")
    assert task_table["rows"][0][5]["text"] == "OH: 08:00 EM: CC EMAIL CLAIM"
    html = service.render_html(report)
    assert "Përshkrimi i gjatë" not in html
    assert "[[added]]" not in html


def test_m1_migration_creates_independent_report_storage():
    from io import StringIO
    from unittest.mock import patch
    from alembic.migration import MigrationContext
    from alembic.operations import Operations
    from tests.test_migration_graph import migration_scripts

    migration = migration_scripts().get_revision("0147_m1_reporting_points").module
    output = StringIO()
    context = MigrationContext.configure(dialect_name="postgresql", opts={"as_sql": True, "output_buffer": output})
    with patch.object(migration, "op", Operations(context)):
        migration.upgrade()
    sql = output.getvalue()
    assert "CREATE TABLE m1_reporting_points_reports" in sql
    assert "manual_answers JSONB" in sql
    assert "auto_sent_at TIMESTAMP WITH TIME ZONE" in sql


def test_generate_reloads_all_automatic_sections_and_preserves_manual_answers(monkeypatch):
    note = SimpleNamespace(id="note", content="Original", created_at=at(DAY, 8), status="OPEN")
    task = SimpleNamespace(id="task", title="Original task", created_at=at(DAY, 8), status="TODO")
    entry = SimpleNamespace(id="absence", category=CommonCategory.absences, entry_date=DAY,
                            created_at=at(DAY, 7), assigned_to_user_id=None, created_by_user_id=None,
                            title="Ada", description="From: 08:00 - To: 12:00")
    rows = [[entry], [note], [task]]

    class Result:
        def __init__(self, items):
            self.items = items

        def scalars(self):
            return self

        def all(self):
            return self.items

    class Db:
        calls = 0

        async def execute(self, statement):
            result = Result(rows[self.calls % 3])
            self.calls += 1
            return result

    async def context(db, tasks):
        return tasks, [], {}, {}

    async def ga_context(db, day, cutoff):
        return {"ga_progress": [], "delivery": []}

    monkeypatch.setattr(service, "prepare_task_context", context)
    monkeypatch.setattr(service, "build_ga_task_data", ga_context)
    answers = {"plan_changes": "Plani aktual", "reorganization": "Po"}
    report = SimpleNamespace(report_date=DAY, data={}, manual_answers=answers, generated_at=None)
    db = Db()

    asyncio.run(service.refresh_report(db, report, at(DAY, 8, 30)))
    assert [row["id"] for row in report.data["notes"]] == ["note"]
    assert report.data["tasks"][0]["status"] == "TODO"

    note.status = "CLOSED"
    rows[1].append(SimpleNamespace(id="new-note", content="Created at 09:45", created_at=at(DAY, 9, 45), status="OPEN"))
    task.status, task.title = "DONE", "Updated task"
    rows[2].append(SimpleNamespace(id="new-task", title="Created at 09:59", created_at=at(DAY, 9, 59), status="TODO"))
    entry.description = "From: 08:00 - To: 16:00"
    rows[0].append(SimpleNamespace(id="delay", category=CommonCategory.delays, entry_date=DAY,
                                  created_at=at(DAY, 9, 30), assigned_to_user_id=None, created_by_user_id=None,
                                  title="Beni", description="Start: 09:00 Until: 10:00"))

    asyncio.run(service.refresh_report(db, report, at(DAY, 10)))
    assert db.calls == 6
    assert report.generated_at == at(DAY, 10)
    assert report.data["captured_until"] == at(DAY, 10).isoformat()
    assert report.manual_answers == answers
    assert [row["id"] for row in report.data["notes"]] == ["new-note"]
    assert report.data["tasks"][0]["status"] == "DONE"
    assert report.data["tasks"][0]["title"] == "Updated task"
    assert report.data["tasks"][1]["id"] == "new-task"
    assert report.data["absences"][0]["to"] == "16:00"
    assert report.data["delays"][0]["id"] == "delay"


def test_attendance_created_after_generation_is_excluded():
    entry = SimpleNamespace(id="future", category=CommonCategory.delays, entry_date=DAY,
                            created_at=at(DAY, 10, 1), assigned_to_user_id=None, created_by_user_id=None,
                            title="Beni", description=None)
    data = service.select_sections([entry], [], [], [], DAY, now=at(DAY, 10))
    assert data["delays"] == []


def test_historical_reports_keep_their_saved_capture(monkeypatch):
    from unittest.mock import AsyncMock
    build = AsyncMock()
    monkeypatch.setattr(service, "build_data", build)
    report = SimpleNamespace(report_date=DAY, data={"notes": [{"id": "old"}]}, generated_at=at(DAY, 10))
    with pytest.raises(ValueError):
        asyncio.run(service.refresh_report(None, report, at(date(2026, 10, 14), 10)))
    build.assert_not_awaited()
    assert report.data == {"notes": [{"id": "old"}]}


def test_exports_show_the_generation_cutoff():
    report = {"report_date": DAY.isoformat(), "data": {"captured_until": at(DAY, 10).isoformat()}, "manual_answers": {}}
    assert "Gjendja deri më: 13.10.2026 10:00" in service.render_plain_text(report)


def rich_export_fixture():
    return {"report_date": DAY.isoformat(), "manual_answers": {"plan_changes": "Plani me ë dhe ç", "reorganization": "Pa ndryshim"},
        "data": {"captured_until": at(DAY, 8, 30).isoformat(),
            "absences": [{"id": "a", "person": "Ada", "from": "08:00", "to": "12:00"}],
            "delays": [{"id": "b", "person": "Beni", "start": "08:00", "until": "08:04"}],
            "notes": [{"id": "n", "source": "PX Notes", "author": "Ada Test",
                "content": "GA: NOTE ! ? F M2/3\n[[added]]E shtuar[[/added]] dhe [[done]]E kryer[[/done]]\n[[added]][[done]]Të dyja[[/done]][[/added]]", "created_at": at(DAY, 8).isoformat()}],
            "tasks": [{"id": "t", "title": "EF: 08:00: ! ? M2/3", "status": "IN_PROGRESS", "created_at": at(DAY, 8).isoformat(),
                "assignees": "EF", "department": "DEV", "am_pm": "AM", "task_type": "1H"}]}}


def test_m1_note_formatting_survives_email_word_png_and_excel(tmp_path):
    import io
    from docx import Document
    from openpyxl import load_workbook
    from PIL import Image
    report = rich_export_fixture()
    html = service.render_html(report)
    assert "background:#BFDBFE" in html
    assert "text-decoration:line-through" in html
    assert "GA: NOTE ! ? F M2/3" in html and "[[added]]" not in html
    attachments = service.report_attachments(report)
    word = Document(io.BytesIO(attachments[0][1]))
    runs = word.tables[2].rows[1].cells[2].paragraphs[0].runs
    assert any(run.text == "E kryer" and run.font.strike for run in runs)
    assert any(run.text == "E shtuar" and 'BFDBFE' in run._element.xml for run in runs)
    workbook = load_workbook(io.BytesIO(attachments[2][1]), rich_text=True)
    assert len(workbook["PAMJA M1"]._images) == 1
    sheet = workbook.active
    note = next(cell.value for row in sheet for cell in row if str(cell.value).startswith("GA: NOTE"))
    assert any(getattr(item, "text", None) == "E kryer" and item.font.strike for item in note)
    assert any(getattr(item, "text", None) == "E shtuar" and item.font.color.rgb == "00172554" for item in note)
    with Image.open(io.BytesIO(attachments[1][1])) as image:
        colors = {color for _, color in image.getcolors(image.width * image.height)}
        assert (191, 219, 254) in colors and (209, 250, 229) in colors
    assert "GA: NOTE ! ? F M2/3" in service.render_plain_text(report)


def test_rich_note_tabs_use_the_same_column_positions_as_png_wrapping():
    from app.services.report_rich_text import marked_note_runs, wrapped_runs
    runs = marked_note_runs("AB[[added]]\tCD[[/added]]\n\nEF")
    lines = ("AB  CD", "", "EF")
    rich_lines = wrapped_runs(runs, lines)
    assert ["".join(run["text"] for run in line) for line in rich_lines] == list(lines)
    assert rich_lines[0][-1]["fill"] == "#BFDBFE"
