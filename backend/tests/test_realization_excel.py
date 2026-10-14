import io
import unittest

from openpyxl import load_workbook

from app.services.realization_calculator import (
    QUESTION_LABELS,
    REPORT_QUESTION_SECTIONS,
    build_live_questions,
)
from app.services.realization_excel import build_realization_workbook


class TestRealizationExcel(unittest.TestCase):
    def test_live_questions_cover_the_full_reference_catalog(self) -> None:
        questions = build_live_questions(
            {
                "weekly_planned_count": 5,
                "weekly_completed_count": 2,
                "weekly_additional_count": 1,
                "weekly_fast_task_count": 1,
                "daily_planned_count": 2,
                "daily_completed_count": 1,
                "counters": {"in_progress_count": 1, "no_progress_count": 0},
                "tasks": [],
                "observations": [],
                "daily_timeline": [],
                "attendance": [],
            }
        )
        expected_keys = [
            key
            for _section_title, question_keys in REPORT_QUESTION_SECTIONS
            for key in question_keys
        ]
        self.assertEqual([question["key"] for question in questions], expected_keys)
        self.assertEqual(len(questions), 17)
        by_key = {question["key"]: question for question in questions}
        for key in (
            "plan_completed",
            "no_progress_tasks",
            "in_progress_tasks",
            "new_tasks_added",
            "closed_tasks",
            "frequent_delays",
            "unexpected_absences",
        ):
            self.assertTrue(by_key[key]["source_status"].startswith("AUTO"), key)
        # Managerial judgment remains manual even when supporting evidence is absent.
        for key in (
            "approved_postponement",
            "respected_meetings",
            "helped_colleague",
            "requested_extra_tasks",
            "gave_proposal",
            "affected_other_plan",
            "repeated_after_clarification",
        ):
            self.assertEqual(by_key[key]["source_status"], "MANUAL_UNANSWERED", key)
        self.assertEqual(by_key["extra_engagement"]["source_status"], "AUTO")
        self.assertFalse(by_key["helped_colleague"]["auto_value"])
        self.assertEqual(by_key["respected_meetings"]["auto_value"]["missed_meeting_evidence"], 0)
        self.assertEqual(by_key["respected_meetings"]["source_status"], "MANUAL_UNANSWERED")

    def test_daily_questions_use_the_selected_days_counts(self) -> None:
        questions = build_live_questions(
            {
                "date": "2026-09-17",
                "daily_planned_count": 4,
                "daily_completed_count": 3,
                "weekly_planned_count": 20,
                "weekly_completed_count": 12,
                "counters": {
                    "additional_count": 2,
                    "fast_task_count": 1,
                    "in_progress_count": 1,
                    "no_progress_count": 0,
                    "tardiness_count": 1,
                },
                "tasks": [],
                "observations": [],
                "attendance": [{"id": "attendance-1", "type": "VONESE"}],
            }
        )
        by_key = {question["key"]: question for question in questions}
        self.assertEqual(by_key["plan_completed"]["auto_value"]["planned"], 4)
        self.assertEqual(by_key["plan_completed"]["auto_value"]["completed"], 3)
        self.assertEqual(by_key["new_tasks_added"]["auto_value"], {
            "yes": True,
            "total": 2,
            "completed": 0,
            "in_progress": 0,
            "todo": 2,
            "postponed": 0,
        })
        self.assertEqual(by_key["frequent_delays"]["auto_value"]["attendance_tardiness"], 1)
        self.assertTrue(by_key["frequent_delays"]["auto_value"]["answer"])
        self.assertEqual(by_key["frequent_delays"]["auto_value"]["threshold"], 1)
        self.assertEqual(by_key["unexpected_absences"]["source_status"], "AUTO")

    def test_export_has_department_evidence_guide_and_weekly_bonus(self) -> None:
        payload = build_realization_workbook(
            week_start="2026-08-03",
            week_end="2026-08-07",
            departments=[
                {
                    "name": "Development",
                    "status": "CALCULATED",
                    "people": [
                        {
                            "user_name": "Test Person",
                            "planned_count": 5,
                            "completed_on_time_count": 4,
                            "completed_late_count": 0,
                            "additional_count": 1,
                            "suggested_level": "A",
                            "suggested_symbol": "+",
                            "final_level": None,
                            "final_symbol": None,
                            "facts_json": {
                                "weekly_progress_percent": 80,
                                "questions": [
                                    {
                                        "key": "plan_completed",
                                        "label": "A janë përfunduar detyrat sipas planit?",
                                        "auto_value": {"answer": False, "planned": 5, "completed": 4},
                                        "final_value": None,
                                        "source_status": "AUTO",
                                        "evidence_ids": ["task-1"],
                                        "explanation": "",
                                    },
                                    {
                                        "key": "respected_meetings",
                                        "label": "A i ka respektuar takimet?",
                                        "auto_value": {"missed_meeting_evidence": 0},
                                        "final_value": None,
                                        "source_status": "MANUAL_ANSWERED",
                                        "evidence_ids": [],
                                        "explanation": "",
                                        "manager_comment": "Nuk kishte takim",
                                    },
                                ],
                                "tasks": [
                                    {
                                        "task_id": "task-1",
                                        "match_key": "task-1",
                                        "title": "Task",
                                        "source_type": "system",
                                        "classification": "completed_on_time",
                                    }
                                ],
                                "observations": [],
                            },
                        }
                    ],
                }
            ],
        )
        workbook = load_workbook(io.BytesIO(payload), data_only=False)
        self.assertEqual(
            workbook.sheetnames,
            ["Përmbledhje", "Development", "Evidenca", "Udhëzuesi", "Udhëzuesi i Vlerësimit", "Komentet ditore", "Diamantat"],
        )
        values = " ".join(
            str(cell.value or "")
            for sheet in workbook.worksheets
            if sheet.title != "Udhëzuesi i Vlerësimit"
            for row in sheet.iter_rows()
            for cell in row
        ).lower()
        # The operational export keeps the weekly bonus and the 17-question guide.
        self.assertIn("bonusi javor", values)
        self.assertIn("17 pyetjet e raportit", values)

        development = workbook["Development"]
        development_values = " ".join(
            str(cell.value or "") for row in development.iter_rows() for cell in row
        ).lower()
        self.assertIn("40", development_values)
        self.assertIn("nuk aplikohet / nuk dihet", development_values)
        self.assertIn("nuk kishte takim", development_values)
        self.assertIn("totali i javës", development_values)
        self.assertIn("nënshkrimet", development_values)

    def test_live_export_is_populated_and_clearly_not_final(self) -> None:
        payload = build_realization_workbook(
            week_start="2026-08-03",
            week_end="2026-08-07",
            departments=[
                {
                    "name": "Development",
                    "status": "AKTUAL (SNAPSHOT DITOR)",
                    "report_mode": "LIVE_DAILY",
                    "people": [
                        {
                            "user_name": "Live Person",
                            "planned_count": 6,
                            "completed_on_time_count": 1,
                            "completed_late_count": 0,
                            "additional_count": 2,
                            "suggested_level": None,
                            "suggested_symbol": None,
                            "final_level": None,
                            "final_symbol": None,
                            "facts_json": {
                                "report_mode": "LIVE_DAILY",
                                "weekly_progress_percent": 16.7,
                                "questions": [
                                    {
                                        "key": "plan_completed",
                                        "label": "A janë përfunduar detyrat sipas planit?",
                                        "auto_value": {"answer": False, "planned": 5, "completed": 0},
                                        "source_status": "AUTO",
                                        "evidence_ids": [],
                                    }
                                ],
                                "daily_timeline": [
                                    {
                                        "date": "2026-08-04",
                                        "daily_progress_percent": 0,
                                        "weekly_progress_percent": 16.7,
                                        "planned_count": 5,
                                        "completed_count": 0,
                                        "weekly_planned_count": 6,
                                        "weekly_completed_count": 1,
                                        "additional_count": 1,
                                        "attendance": [],
                                    }
                                ],
                            },
                        }
                    ],
                }
            ],
        )
        workbook = load_workbook(io.BytesIO(payload), data_only=False)
        summary = workbook["Përmbledhje"]
        development = workbook["Development"]

        self.assertEqual(summary["B5"].value, "Live Person")
        self.assertEqual(summary["C5"].value, 6)
        self.assertEqual(summary["D5"].value, 1)
        self.assertEqual(summary["H5"].value, "Në pritje të FINAL")
        self.assertIn("AKTUAL NGA SNAPSHOT-ET DITORE", summary["A2"].value)
        values = " ".join(
            str(cell.value or "")
            for row in development.iter_rows()
            for cell in row
        )
        self.assertIn("Sot: 0/5 (0%)", values)
        self.assertIn("Vlerësimi final", values)
        self.assertIn("PËR KONFIRMIM", values)
        for section_title, question_keys in REPORT_QUESTION_SECTIONS:
            self.assertIn(section_title, values)
            for key in question_keys:
                self.assertIn(QUESTION_LABELS[key], values)


if __name__ == "__main__":
    unittest.main()


def test_export_uses_daily_percent_and_dated_comments_and_confirmed_letter():
    person = {"user_name": "Rinesa", "weekly_manager_review": {"level": "B", "created_by_name": "Manager", "comment": "Reviewed"},
              "facts_json": {"weekly_progress_percent": 84.8,
                             "weekly_evaluation": {"level": "C", "answers_complete": True, "reasons": ["Explanation"]},
                             "daily_timeline": [{"date": "2026-10-05", "daily_progress_percent": 79.2,
                                                 "weekly_progress_percent": 20, "person_comment": {"comment": "=not a formula", "author": "Manager"}},
                                                {"date": "2026-10-06", "on_leave": True, "daily_progress_percent": None}]}}
    book = load_workbook(io.BytesIO(build_realization_workbook(week_start="2026-10-05", week_end="2026-10-09",
        departments=[{"name": "DEV", "report_mode": "LIVE_DAILY", "people": [person]}])))
    rows = {row[0].value: row for row in book["DEV"].iter_rows() if row[0].value}
    assert rows["Snapshot ditor — 2026-10-05"][1].value == 0.792
    assert rows["Snapshot ditor — 2026-10-06"][1].value == "Pushim"
    assert rows["Propozimi për nivelin e vlerësimit"][1].value == "C"
    assert rows["Vlerësimi final"][1].value == "B"
    assert book["Komentet ditore"]["C5"].value == "2026-10-05"
    assert book["Komentet ditore"]["F5"].value == "=not a formula"
    assert book["Komentet ditore"]["F5"].data_type == "s"
    assert book["Përmbledhje"]["H5"].value == "B"


def test_weekly_export_preserves_all_daily_comment_sources_and_dates():
    days = [
        {"date": "2026-10-12", "daily_progress_percent": 50,
         "person_comment": {"comment": "Komenti i së hënës\nRreshti i dytë", "author": "Përgjegjësi"},
         "manual_answers": {"gave_proposal": {"value": True, "comment": "Propozoi automatizim", "answered_by_name": "Manager"}},
         "close_event": {"daily_comment": "Mbyllja e së hënës"},
         "tasks": [{"title": "Detyra A", "daily_report_comment": "Detyra e së hënës"}]},
        {"date": "2026-10-13", "daily_progress_percent": 80,
         "person_comment": {"comment": "Komenti i së martës", "author": "Përgjegjësi"},
         "manual_answers": {"helped_colleague": {"value": True, "comment": "Ndihmoi kolegun", "answered_by_name": "Manager"}},
         "close_event": {"daily_comment": "Mbyllja e së martës"},
         "tasks": [{"title": "Detyra A", "daily_report_comment": "=Komenti i detyrës të martën"}]},
    ]
    person = {"user_name": "Elsa", "facts_json": {"daily_timeline": days}}
    book = load_workbook(io.BytesIO(build_realization_workbook(
        week_start="2026-10-12", week_end="2026-10-16",
        departments=[{"name": "DEV", "report_mode": "LIVE_DAILY", "people": [person]}],
    )), data_only=False)
    entries = list(book["Komentet ditore"].iter_rows(min_row=5, values_only=True))
    assert len(entries) == 8
    by_comment = {entry[5]: entry for entry in entries}
    expected = [
        ("Komenti i së hënës\nRreshti i dytë", "2026-10-12", "Komenti ditor", "Përgjegjësi"),
        ("Propozoi automatizim", "2026-10-12", QUESTION_LABELS["gave_proposal"], "Manager"),
        ("Mbyllja e së hënës", "2026-10-12", "Mbyllja ditore", None),
        ("Detyra e së hënës", "2026-10-12", "Detyra A", None),
        ("Komenti i së martës", "2026-10-13", "Komenti ditor", "Përgjegjësi"),
        ("Ndihmoi kolegun", "2026-10-13", QUESTION_LABELS["helped_colleague"], "Manager"),
        ("Mbyllja e së martës", "2026-10-13", "Mbyllja ditore", None),
        ("=Komenti i detyrës të martën", "2026-10-13", "Detyra A", None),
    ]
    person_sheet = "\n".join(str(cell.value or "") for row in book["DEV"].iter_rows() for cell in row)
    for comment, day, source, author in expected:
        assert by_comment[comment] == ("DEV", "Elsa", day, source, author, comment)
        assert comment in person_sheet
    formula_comment = next(cell for row in book["Komentet ditore"].iter_rows() for cell in row
                           if cell.value == "=Komenti i detyrës të martën")
    assert formula_comment.data_type == "s"


def test_excel_exports_task_and_proposal_diamonds_in_totals_and_details():
    task_award = {"task_id": "task-1", "task_title": "[[b]]Raporti[[/b]]", "source": "TASK",
                  "completed_at": "2026-10-12T12:00:00+02:00",
                  "review": {"id": "review-1", "diamond_score": 1, "comment": "Punë e mirë", "reviewer_name": "Manager"}}
    proposal = {"task_id": None, "realization_period_id": "period-1", "task_title": "Propozim për përmirësim",
                "source": "REALIZATION_PROPOSAL", "completed_at": "2026-10-13T12:00:00+02:00",
                "review": {"id": "result-1", "diamond_score": 1, "comment": "=Propozimi\nKursen 20 minuta", "reviewer_name": "Admin"}}
    people = [
        {"user_name": "Elsa", "facts_json": {"weekly_diamond_count": 2, "weekly_task_diamond_count": 1,
                                              "weekly_proposal_diamond_count": 1, "weekly_diamonds": [task_award, proposal]}},
        {"user_name": "Rinesa", "facts_json": {}},
    ]
    for report_mode in ["LIVE_DAILY", "FINAL_WEEKLY"]:
        book = load_workbook(io.BytesIO(build_realization_workbook(
            week_start="2026-10-12", week_end="2026-10-16",
            departments=[{"name": "DEV", "report_mode": report_mode, "people": people}],
        )), data_only=False)
        assert [book["Përmbledhje"].cell(4, col).value for col in [11, 12, 13]] == [
            "Diamanta gjithsej", "Diamanta nga taskat", "Diamanta nga propozimet"]
        assert [book["Përmbledhje"].cell(5, col).value for col in [11, 12, 13]] == [2, 1, 1]
        assert [book["Përmbledhje"].cell(6, col).value for col in [11, 12, 13]] == [0, 0, 0]
        assert book["Përmbledhje"].auto_filter.ref == "A4:M6"
        rows = {row[0].value: row for row in book["DEV"].iter_rows() if row[0].value}
        assert rows["Diamanta gjithsej"][1].value == 2
        assert rows["Diamanta nga taskat"][1].value == 1
        assert rows["Diamanta nga propozimet"][1].value == 1
        assert "Punë e mirë" in rows["Diamanta gjithsej"][2].value
        assert "Kursen 20 minuta" in rows["Diamanta gjithsej"][2].value
        assert rows["Diamanta gjithsej"][3].value == 0
        details = list(book["Diamantat"].iter_rows(min_row=5, values_only=True))
        assert len(details) == 2
        assert details[0] == ("DEV", "Elsa", "2026-10-12", "Taskë e vlerësuar", "Raporti", 1, "Punë e mirë", "Manager", "task-1")
        assert details[1] == ("DEV", "Elsa", "2026-10-13", "Propozim për përmirësim", "Propozim për përmirësim", 1,
                              "=Propozimi\nKursen 20 minuta", "Admin", "period-1")
        assert book["Diamantat"]["G6"].data_type == "s"
        assert book["Diamantat"]["G6"].alignment.wrap_text
        assert sum(entry[5] for entry in details) == sum(book["Përmbledhje"].cell(row, 11).value for row in [5, 6])


def test_snapshot_argument_highlights_comments_before_metrics_and_grows_to_fit():
    comment = "Koment i rëndësishëm për këtë ditë\n" + "Sqarim i plotë për progresin. " * 20
    person = {"user_name": "Elsa", "facts_json": {"daily_timeline": [
        {"date": "2026-10-12", "daily_progress_percent": 75,
         "person_comment": {"comment": comment, "author": "Admin"}},
        {"date": "2026-10-13", "on_leave": True, "daily_progress_percent": None},
    ]}}
    book = load_workbook(io.BytesIO(build_realization_workbook(
        week_start="2026-10-12", week_end="2026-10-16",
        departments=[{"name": "DEV", "people": [person]}],
    )), rich_text=True)
    rows = {row[0].value: row for row in book["DEV"].iter_rows() if row[0].value}
    highlighted = rows["Snapshot ditor — 2026-10-12"][2]
    assert str(highlighted.value).startswith("KOMENTET E DITËS\n")
    assert str(highlighted.value).index(comment) < str(highlighted.value).index("Java deri atë ditë")
    assert highlighted.value[0].font.b is True
    assert highlighted.value[0].font.color.rgb == "00075985"
    assert highlighted.value[0].font.sz == 12
    assert highlighted.value[1].font.sz == 10
    assert highlighted.fill.fgColor.rgb == "00CFFAFE"
    assert highlighted.alignment.wrap_text
    assert book["DEV"].row_dimensions[highlighted.row].height > 100
    leave = rows["Snapshot ditor — 2026-10-13"][2]
    assert "None%" not in str(leave.value)
    assert "KOMENTET E DITËS" not in str(leave.value)
    assert len(book["Komentet ditore"]["F5"].value) == len(comment)


def test_diamond_detail_date_uses_application_timezone():
    from unittest.mock import patch
    from app.services.realization_excel import _diamond_date

    with patch("app.services.realization_excel.settings.APP_TIMEZONE", "Europe/Budapest"):
        assert _diamond_date("2026-10-12T22:30:00Z") == "2026-10-13"


def test_export_without_awards_keeps_an_empty_diamond_details_sheet():
    book = load_workbook(io.BytesIO(build_realization_workbook(
        week_start="2026-10-12", week_end="2026-10-16", departments=[],
    )))
    assert book["Diamantat"].max_row == 4
    assert book["Diamantat"].auto_filter.ref == "A4:I4"
    assert book["Përmbledhje"].max_column == 13
