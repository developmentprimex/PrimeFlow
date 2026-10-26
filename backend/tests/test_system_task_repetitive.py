import io
import unittest
import uuid
from datetime import datetime, time, timezone
from types import SimpleNamespace
from unittest.mock import AsyncMock, Mock, patch

from fastapi import FastAPI
from fastapi.testclient import TestClient
from openpyxl import load_workbook

from app.db import get_db
from app.api.routers.exports import (
    export_system_task_templates_xlsx,
    export_system_tasks_xlsx,
    preview_system_tasks,
    router as exports_router,
)
from app.api.routers.system_tasks import create_system_task_template, update_system_task_template
from app.models.enums import CommonApprovalStatus, FrequencyType, SystemTaskScope, UserRole
from app.models.system_task_template import SystemTaskTemplate
from app.schemas.system_task_template import SystemTaskTemplateCreate, SystemTaskTemplateUpdate


class _Result:
    def __init__(self, rows):
        self.rows = rows

    def scalars(self):
        return self

    def all(self):
        return self.rows

    def scalar_one_or_none(self):
        return self.rows[0] if self.rows else None


def _db(*results):
    return SimpleNamespace(
        execute=AsyncMock(side_effect=[_Result(rows) for rows in results]),
        add=Mock(),
        flush=AsyncMock(),
        commit=AsyncMock(),
        refresh=AsyncMock(),
    )


def _template(**overrides):
    values = dict(
        id=uuid.uuid4(), title="System task", scope=SystemTaskScope.ALL.value,
        frequency=FrequencyType.DAILY.value, timezone="Europe/Budapest",
        due_time=time(9), lookahead=14, interval=1, duration_days=1,
        is_active=True, requires_alignment=False, show_in_weekly_planner=False,
        approval_status=CommonApprovalStatus.approved,
        created_at=datetime(2026, 10, 26, tzinfo=timezone.utc),
    )
    values.update(overrides)
    return SystemTaskTemplate(**values)


class TestSystemTaskRepetitive(unittest.IsolatedAsyncioTestCase):
    def setUp(self):
        self.user = SimpleNamespace(id=uuid.uuid4(), role=UserRole.ADMIN, full_name="Test User", username="test")

    def test_database_defaults_existing_tasks_to_no(self):
        column = SystemTaskTemplate.__table__.c.is_repetitive
        self.assertFalse(column.nullable)
        self.assertEqual(str(column.server_default.arg).lower(), "false")

    async def test_create_saves_explicit_flag_and_defaults_to_no(self):
        for value in (True, False, None):
            with self.subTest(value=value):
                payload = SystemTaskTemplateCreate(
                    title="System task", frequency=FrequencyType.DAILY, is_repetitive=value,
                )
                db = _db()
                with (
                    patch("app.api.routers.system_tasks._validate_replacement_users", new=AsyncMock()),
                    patch("app.api.routers.system_tasks._sync_template_slots_from_payload", new=AsyncMock()),
                    patch("app.api.routers.system_tasks._template_definition_to_out", new=AsyncMock()),
                ):
                    await create_system_task_template(payload, db=db, user=self.user)
                template = db.add.call_args.args[0]
                self.assertEqual(template.is_repetitive, bool(value))
                self.assertEqual(template.frequency, FrequencyType.DAILY.value)
                db.commit.assert_awaited_once()

    async def test_update_round_trips_flag_without_rescheduling(self):
        template = _template(is_repetitive=False)
        for payload, expected in (
            (SystemTaskTemplateUpdate(is_repetitive=True), True),
            (SystemTaskTemplateUpdate(description="Updated"), True),
            (SystemTaskTemplateUpdate(is_repetitive=False), False),
        ):
            with self.subTest(payload=payload.model_dump(exclude_unset=True)):
                # Template lookup, then alignment roles/users and assignee slots for the response.
                db = _db([template], [], [], [])
                with (
                    patch("app.api.routers.system_tasks._validate_replacement_users", new=AsyncMock()),
                    patch("app.api.routers.system_tasks._reset_template_slots_next_run_at", new=AsyncMock()) as reset,
                    patch("app.api.routers.system_tasks._regenerate_template_after_assignment_change", new=AsyncMock()) as regenerate,
                ):
                    result = await update_system_task_template(template.id, payload, db=db, user=self.user)
                self.assertEqual(template.is_repetitive, expected)
                self.assertEqual(result.is_repetitive, expected)
                self.assertEqual(template.frequency, FrequencyType.DAILY.value)
                reset.assert_not_awaited()
                regenerate.assert_not_awaited()
                db.commit.assert_awaited_once()

    async def test_excel_exports_show_po_jo_and_include_column_in_filter(self):
        templates = [_template(is_repetitive=True), _template(is_repetitive=False), _template()]
        tasks = [SimpleNamespace(
            id=uuid.uuid4(), assigned_to=None, created_at=template.created_at,
            title=template.title, description=None, finish_period=None,
        ) for template in templates]
        for export in (export_system_task_templates_xlsx, export_system_tasks_xlsx):
            with self.subTest(export=export.__name__):
                if export is export_system_task_templates_xlsx:
                    db = _db(templates, [])
                else:
                    db = _db(templates, list(zip(tasks, templates)), [], [], [])
                response = await export(db=db, user=self.user)
                content = b"".join([chunk async for chunk in response.body_iterator])
                sheet = load_workbook(io.BytesIO(content)).active
                columns = {cell.value: cell.column for cell in sheet[5]}
                column = columns["Detyrë e përsëritur"]
                self.assertEqual([sheet.cell(row, column).value for row in (6, 7, 8)], ["Po", "Jo", "Jo"])
                self.assertEqual(column, sheet.max_column)
                self.assertTrue(sheet.cell(6, column).alignment.wrap_text)
                self.assertTrue(sheet.auto_filter.ref.endswith("8"))

    async def test_preview_shows_same_flag_as_excel(self):
        template = _template(is_repetitive=True)
        task = SimpleNamespace(
            id=uuid.uuid4(), assigned_to=None, created_at=template.created_at,
            title=template.title, description=None, finish_period=None,
        )
        result = await preview_system_tasks(
            db=_db([template], [(task, template)], [], [], []), user=self.user,
        )
        column = result.headers.index("Detyrë e përsëritur")
        self.assertEqual(result.rows[0][column], "Po")

    async def test_template_excel_export_allows_staff_manager_and_admin(self):
        for role in UserRole:
            with self.subTest(role=role):
                user = SimpleNamespace(id=uuid.uuid4(), role=role, full_name="Test User", username="test")
                template = _template(title="Export accessible task", is_repetitive=True)
                response = await export_system_task_templates_xlsx(
                    mode="all", db=_db([template], []), user=user,
                )
                self.assertEqual(response.status_code, 200)
                content = b"".join([chunk async for chunk in response.body_iterator])
                workbook = load_workbook(io.BytesIO(content))
                self.assertEqual(workbook.active.cell(6, 7).value, template.title)

    def test_template_excel_export_still_requires_authentication(self):
        app = FastAPI()
        app.include_router(exports_router, prefix="/exports")
        db = _db()
        app.dependency_overrides[get_db] = lambda: db
        with TestClient(app) as client:
            response = client.get("/exports/system-task-templates.xlsx?mode=all")
        self.assertEqual(response.status_code, 401)
        db.execute.assert_not_awaited()


if __name__ == "__main__":
    unittest.main()
