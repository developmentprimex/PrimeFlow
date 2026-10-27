import io
import unittest
from datetime import date, time
from types import SimpleNamespace
from unittest.mock import AsyncMock, MagicMock, patch

from openpyxl import load_workbook

from app.api.routers import exports
from app.models.enums import UserRole


class GaTimeExportFilterTests(unittest.IsolatedAsyncioTestCase):
    async def export(self, selected_date=None, role=UserRole.ADMIN):
        templates = [
            SimpleNamespace(day_of_week=day, start_time=time(8), content=f"Day {day}")
            for day in range(7)
        ]
        imported = [SimpleNamespace(day_date=date(2026, 10, 28), start_time=time(8), content="Wednesday import")]
        results = []
        for entries in (templates, imported):
            result = MagicMock()
            result.scalars.return_value.all.return_value = entries
            results.append(result)
        db = SimpleNamespace(execute=AsyncMock(side_effect=results))
        rows = [SimpleNamespace(
            nr_label="1", label="08:00 - 09:00", start_time=time(8), end_time=time(9),
            comment="Start comment", comments=[{"column": "end", "content": "End comment"}],
        )]
        user = SimpleNamespace(role=role, full_name="Test User", username="test-user")
        with patch.object(exports, "_resolve_ga_user", AsyncMock(return_value=SimpleNamespace(id="ga"))), \
             patch.object(exports, "get_ga_time_table_rows", AsyncMock(return_value=rows)):
            response = await exports.export_ga_time_xlsx(
                week_start=date(2026, 10, 26), selected_date=selected_date, db=db, user=user,
            )
        payload = b"".join([chunk async for chunk in response.body_iterator])
        return load_workbook(io.BytesIO(payload)).active, response

    async def test_every_user_role_can_export_day_and_week(self):
        for role in UserRole:
            for selected in (None, date(2026, 10, 28)):
                with self.subTest(role=role, selected=selected):
                    ws, response = await self.export(selected, role=role)
                    self.assertEqual(response.status_code, 200)
                    self.assertEqual(ws.max_column, 4 if selected else 9)

    async def test_selected_wednesday_exports_four_columns_and_correct_entries(self):
        ws, response = await self.export(date(2026, 10, 28))
        self.assertEqual(ws.max_column, 4)
        self.assertEqual([cell.value for cell in ws[4]][:3], ["NR", "TIME", "KOMENT"])
        self.assertIn("28.10.2026", ws["D4"].value)
        self.assertEqual(ws["D5"].value, "Day 2\nWednesday import")
        self.assertEqual(ws["C5"].value, "Start comment")
        self.assertEqual(ws["A1"].value, "GA TIME TABLE (28.10.2026)")
        self.assertIn("28_10_26", response.headers["content-disposition"])
        self.assertEqual(ws.column_dimensions["D"].width, 40)

    async def test_selected_friday_uses_friday_instead_of_monday(self):
        ws, _ = await self.export(date(2026, 10, 30))
        self.assertEqual(ws.max_column, 4)
        self.assertEqual(ws["D5"].value, "Day 4")

    async def test_unfiltered_export_keeps_week_and_end_comments(self):
        ws, _ = await self.export()
        self.assertEqual(ws.max_column, 9)
        self.assertEqual(ws["D5"].value, "Day 0")
        self.assertEqual(ws["H5"].value, "Day 4")
        self.assertEqual(ws["I4"].value, "KOMENT")
        self.assertEqual(ws["I5"].value, "End comment")
        self.assertEqual(ws["A1"].value, "GA TIME TABLE (26.10.2026 - 30.10.2026)")
